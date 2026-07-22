#!/usr/bin/env python3
"""Generate client-facing Word document from ENRICHMENT_CLIENT.md content."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "ENRICHMENT_CLIENT.docx"


def set_cell_shading(cell, fill: str) -> None:
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {qn("w:fill"): fill, qn("w:val"): "clear"})
    shading.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "1F4E79") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(hdr_cells[i], header_fill)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, text in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = text
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading("Enriquecimiento de datos empresariales", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Guía para el cliente")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.size = Pt(13)
    doc.add_paragraph(
        "Este documento explica cómo identificamos y consolidamos la información pública de cada "
        "organización a partir de su nombre legal o comercial, qué campos devolvemos y cómo "
        "interpretar el grado de confianza de cada resultado."
    )

    doc.add_heading("¿Qué hacemos?", level=1)
    doc.add_paragraph(
        "Por cada empresa de entrada generamos un único perfil enriquecido que puede incluir:"
    )
    add_table(
        doc,
        ["Campo", "Descripción"],
        [
            ["Nombre legal", "El nombre con el que se nos proporcionó la organización"],
            ["Nombre comercial", "Denominación pública con la que opera en el mercado"],
            ["LinkedIn", "URL del perfil corporativo en LinkedIn"],
            ["Sitio web", "Página oficial de la empresa (solo dominio corporativo)"],
            ["Dominio", "Dominio principal del sitio web"],
            ["Sector", "Actividad principal (ej. Seguros, Telecomunicaciones)"],
            ["Empleados / seguidores", "Tamaño aproximado según perfil público"],
            ["Teléfono", "Contacto publicado cuando está disponible"],
            ["Sede", "Ubicación de la oficina principal en texto legible"],
            ["Grado de match", "Nivel de confianza en la identificación"],
            ["Puntuación", "Valor numérico de 0 a 100 que resume la fuerza de las evidencias"],
        ],
    )
    doc.add_paragraph(
        "No se devuelven múltiples candidatos por empresa: el sistema elige la mejor identidad "
        "posible según las reglas descritas en este documento."
    )

    doc.add_heading("¿De dónde sale la información?", level=1)
    doc.add_paragraph(
        "Utilizamos varias fuentes complementarias y las cruzamos de forma automática. "
        "No dependemos de una sola fuente para decidir."
    )
    doc.add_paragraph("Flujo resumido:", style="List Bullet")
    doc.add_paragraph("Nombre de la empresa → búsqueda web y perfil LinkedIn", style="List Bullet")
    doc.add_paragraph("Búsqueda web → señales de directorios, sitio web oficial y validación de marca", style="List Bullet")
    doc.add_paragraph("Perfil LinkedIn → datos estructurados de empresa", style="List Bullet")
    doc.add_paragraph("Todas las señales → motor de puntuación → perfil consolidado", style="List Bullet")

    doc.add_heading("1. Búsqueda web", level=2)
    doc.add_paragraph("Buscamos en internet referencias a la empresa para localizar:")
    for item in [
        "Su perfil en LinkedIn",
        "Su sitio web oficial",
        "Menciones en directorios empresariales, registros mercantiles y otras páginas informativas",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Los directorios no se consideran sitio web de la empresa. Solo aportan pistas "
        "(sector, ciudad, a veces la URL real citada en el texto del listado)."
    )

    doc.add_heading("2. Fuente profesional de LinkedIn", level=2)
    doc.add_paragraph(
        "Cuando tenemos un perfil corporativo candidato, lo consultamos para obtener datos "
        "estructurados: nombre comercial, sector, empleados, sede, descripción, etc."
    )

    doc.add_heading("3. Validación del sitio web", level=2)
    doc.add_paragraph(
        "Si encontramos una web candidata, comprobamos que el contenido menciona a la empresa "
        "buscada (título, textos institucionales, aviso legal, «quiénes somos»). Esto evita "
        "asignar páginas de terceros con nombres parecidos."
    )

    doc.add_heading("4. Mapas y fichas locales (último recurso)", level=2)
    doc.add_paragraph(
        "Solo si la búsqueda web no ofrece un sitio fiable, consultamos fichas de negocio en "
        "mapas, priorizando la ciudad o provincia conocida de la empresa."
    )

    doc.add_heading("Cómo construimos el perfil", level=1)
    doc.add_paragraph(
        "No tomamos el primer resultado que aparece. Aplicamos un sistema de puntuación que "
        "valora cada señal y elige la combinación más fiable."
    )
    doc.add_heading("Señales que recogemos", level=2)
    add_table(
        doc,
        ["Origen", "Qué aporta"],
        [
            ["Nombre", "Similitud entre el nombre legal y la marca encontrada"],
            ["LinkedIn", "Coherencia del perfil, posición en resultados, datos de empresa"],
            ["Sitio web", "Coincidencia del dominio con la marca, páginas corporativas"],
            ["Directorios", "Sector, localización, URL citada en el snippet"],
            ["Geografía", "Ciudad, provincia y país del lote vs sede y dominio"],
            ["Tamaño", "Empleados y seguidores cuando están publicados"],
        ],
    )

    doc.add_heading("Reglas de calidad", level=2)
    rules = [
        (
            "Mejor vacío que incorrecto",
            "Si la única web encontrada pertenece a otra empresa del mismo nombre en otro país, "
            "o a un directorio, no la asignamos.",
        ),
        (
            "La web oficial pesa más que un listado",
            "Un dominio corporativo validado tiene prioridad sobre enlaces indirectos.",
        ),
        (
            "Coherencia geográfica",
            "Para lotes de un país concreto, penalizamos dominios claramente asociados a otro mercado.",
        ),
        (
            "Marcas relacionadas ≠ misma entidad",
            "Una gestoría o marca comercial vecina puede aparecer como coincidencia probable, "
            "no como confirmada.",
        ),
        (
            "Portales de terceros",
            "Webs de aseguradoras, consultoras o marketplaces no se asignan como web oficial "
            "salvo que la empresa sea precisamente esa entidad.",
        ),
        (
            "Un solo resultado por empresa",
            "El motor elige la identidad con mayor puntuación global, no una lista de alternativas.",
        ),
    ]
    for n, (title_text, body) in enumerate(rules, start=1):
        p = doc.add_paragraph()
        p.add_run(f"{n}. {title_text}. ").bold = True
        p.add_run(body)

    doc.add_heading("Grados de match", level=1)
    doc.add_paragraph(
        "Cada fila incluye un grado de match que indica cuánta confianza tiene el sistema "
        "en haber identificado a la organización correcta."
    )
    add_table(
        doc,
        ["Grado", "Significado", "Uso recomendado"],
        [
            ["confirmed", "Coincidencia muy sólida: nombre, web y perfil profesional alineados", "Uso directo en CRM, campañas y reporting"],
            ["high_confidence", "Alta confianza; pequeñas diferencias de nombre o datos secundarios", "Uso operativo; revisión puntual si el caso es sensible"],
            ["probable", "Coincidencia razonable; puede ser marca comercial o entidad relacionada", "Revisar antes de decisiones críticas"],
            ["ambiguous", "Señales contradictorias o varios candidatos cercanos", "Revisión manual recomendada"],
            ["partial", "Sitio web identificado pero no un perfil LinkedIn fiable", "Útil para contacto web; completar LinkedIn a mano"],
            ["not_found", "No hay identidad suficientemente fiable", "No usar; puede reintentarse con más contexto"],
        ],
    )

    doc.add_heading("Puntuación numérica (0–100)", level=2)
    add_table(
        doc,
        ["Rango", "Interpretación"],
        [
            ["90–100", "Muy alta confianza → confirmed"],
            ["78–89", "Alta confianza → high_confidence"],
            ["60–77", "Confianza media → probable"],
            ["40–59", "Baja confianza → ambiguous"],
            ["0–39", "Insuficiente → not_found o partial si hay web"],
        ],
    )
    doc.add_paragraph(
        "Si dos candidatos quedan muy próximos en puntuación, el sistema baja automáticamente "
        "el grado de match para reflejar la duda."
    )

    doc.add_heading("Tipos de relación", level=1)
    add_table(
        doc,
        ["Relación", "Significado"],
        [
            ["Misma entidad", "La organización encontrada corresponde a la buscada"],
            ["Marca comercial", "Marca pública relacionada pero no idéntica al nombre legal"],
            ["Matriz / grupo", "Página de la empresa madre o grupo internacional"],
            ["Sucursal", "Oficina o filial geográfica"],
            ["Desconocida", "No hay clasificación clara"],
        ],
    )

    doc.add_heading("¿Cuándo puede faltar información?", level=1)
    add_table(
        doc,
        ["Campo vacío", "Motivo habitual"],
        [
            ["Sitio web", "No hay dominio corporativo fiable en fuentes públicas"],
            ["LinkedIn", "No hay perfil claro o solo aparecen perfiles no corporativos"],
            ["Sede", "El perfil público no publica ubicación"],
            ["Teléfono / empleados", "No están disponibles en las fuentes consultadas"],
            ["Sector", "Perfil sin actividad declarada"],
        ],
    )
    doc.add_paragraph(
        "En estos casos preferimos dejar el campo vacío antes que rellenarlo con datos de "
        "otra empresa o de un portal ajeno."
    )

    doc.add_heading("Ejemplos de interpretación", level=1)
    examples = [
        ("Resultado ideal — confirmed", [
            "Nombre legal: Correduría de Seguros Ejemplo S.L.",
            "Web: ejemploseguros.es",
            "LinkedIn: perfil corporativo con el mismo nombre comercial",
            "Sector: Seguros",
            "→ Identidad lista para usar.",
        ]),
        ("Web sin LinkedIn — partial", [
            "Web oficial encontrada y validada",
            "No hay perfil LinkedIn corporativo claro",
            "→ Útil para web y dominio; completar LinkedIn manualmente si hace falta.",
        ]),
        ("Marca vecina — probable", [
            "Nombre legal muy formal",
            "LinkedIn de una gestoría o marca comercial del mismo sector y zona",
            "→ Revisar si es la entidad operativa que buscáis.",
        ]),
        ("Sin resultado — not_found", [
            "Solo aparecen directorios, noticias o empresas homónimas en otros países",
            "→ No publicamos datos; conviene aportar ciudad o más contexto.",
        ]),
    ]
    for heading, bullets in examples:
        doc.add_heading(heading, level=2)
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("Qué necesitamos para mejores resultados", level=1)
    add_table(
        doc,
        ["Dato de entrada", "Beneficio"],
        [
            ["Nombre legal completo", "Base de todas las búsquedas"],
            ["Ciudad o provincia", "Desambiguación y mejor web/LinkedIn"],
            ["País del lote", "Criterios geográficos coherentes"],
            ["CIF / tax ID", "Futura validación cruzada (cuando esté disponible)"],
        ],
    )

    doc.add_heading("Resumen", level=1)
    for item in [
        "Varias fuentes, cruzadas con scoring — no un solo dato aislado.",
        "Directorios informan, pero no sustituyen a la web oficial.",
        "LinkedIn aporta el perfil estructurado cuando existe match fiable.",
        "Cada fila lleva un grado de match para saber cuándo usar el dato y cuándo revisarlo.",
        "Vacío es mejor que incorrecto — especialmente en webs y perfiles de otras geografías.",
    ]:
        doc.add_paragraph(item, style="List Number")

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Documento de cliente — uso externo")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
